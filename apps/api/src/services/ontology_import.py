"""
Import the world ontology workbook (`data/world-ontology.xlsx`).

`parse_workbook` reads and validates the file and returns the entities; it never
touches the database and never writes to the workbook. `export_json` writes the
generated `data/ontology/world-ontology.json`, and `load_ontology` puts the
entities in `app.ontology_entities`. The Arabic text of every cell is kept exactly
as the workbook has it; only list columns are split, and the search forms that sit
beside the text are made by `src.arabic`.

A broken structure fails with `OntologyImportError`, which lists every problem it
found with the row and the column, so the person who edits the workbook can fix
them in one pass.
"""

from __future__ import annotations

import hashlib
import json
import zipfile
from collections import Counter
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from io import BytesIO
from itertools import batched
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import Boolean, delete, func, literal_column
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import AsyncSession

from src.arabic import normalize_arabic, split_list
from src.models.ontology import OntologyEntity
from src.services.ontology_constraints import ConstraintKind, classify_constraint

SHEET_NAME = "الكيان ومفاهيمه"
# Field name -> the header the workbook has for it. The columns are found by
# these names, so their order in the sheet does not matter.
HEADERS = {
    "label_ar": "الكيان أو الشيء",
    "related_objects": "الأشياء والموضوعات المتصلة",
    "actions_and_uses": "الأفعال والاستعمالات",
    "contextual_concepts": "المفاهيم بحسب السياق",
    "special_constraint": "قيد خاص",
    "domain": "المجال",
    "id": "المعرّف",
}
LIST_FIELDS = ("related_objects", "actions_and_uses", "contextual_concepts")
# Every field but the constraint must be filled in.
REQUIRED_FIELDS = tuple(name for name in HEADERS if name != "special_constraint")

EXPECTED_COUNT = 1000
# The header is searched for in the first rows: the sheet has title rows above it.
HEADER_SEARCH_ROWS = 30
MAX_PROBLEMS_SHOWN = 25
MAX_IDS_SHOWN = 10
# 13 columns per row and asyncpg's limit of 32,767 parameters leave room for 2,500.
ROWS_PER_INSERT = 250


class OntologyImportError(Exception):
    """The workbook is not a valid ontology; `problems` lists what is wrong, one line each."""

    def __init__(self, problems: Sequence[str]) -> None:
        self.problems = tuple(problems)
        shown = self.problems[:MAX_PROBLEMS_SHOWN]
        lines = [f"The ontology workbook is not valid ({len(self.problems)} problem(s)):"]
        lines.extend(f"  - {problem}" for problem in shown)
        if len(self.problems) > len(shown):
            lines.append(f"  - ... and {len(self.problems) - len(shown)} more")
        super().__init__("\n".join(lines))


@dataclass(frozen=True, slots=True)
class OntologyRow:
    """One entity as the workbook has it; `raw` holds its cells untouched."""

    id: str
    label_ar: str
    related_objects: tuple[str, ...]
    actions_and_uses: tuple[str, ...]
    contextual_concepts: tuple[str, ...]
    special_constraint: str | None
    domain: str
    is_catch_all: bool
    raw: dict[str, Any]


@dataclass(frozen=True, slots=True)
class ParsedOntology:
    """A validated workbook: where it came from and every entity in it."""

    source_name: str
    source_sha256: str
    sheet: str
    rows: tuple[OntologyRow, ...]
    warnings: tuple[str, ...] = field(default=())

    @property
    def domains(self) -> list[str]:
        return list(dict.fromkeys(row.domain for row in self.rows))

    @property
    def constraint_counts(self) -> dict[str, int]:
        counts = Counter(row.special_constraint for row in self.rows if row.special_constraint)
        return dict(counts.most_common())


@dataclass(frozen=True, slots=True)
class LoadResult:
    """What a load did to the table."""

    inserted: int
    updated: int
    removed: int


# ─── The file, the sheet and the header ───


def _open_workbook(data: bytes) -> Any:
    try:
        return load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as error:
        message = f"The file is not a readable .xlsx workbook ({type(error).__name__})."
        raise OntologyImportError([message]) from None


def _read_sheet(data: bytes) -> list[tuple[Any, ...]]:
    """Return every row of the ontology sheet, blank rows included, so list positions are row numbers."""
    workbook = _open_workbook(data)
    try:
        wanted = normalize_arabic(SHEET_NAME)
        for name in workbook.sheetnames:
            if normalize_arabic(name) == wanted:
                return [tuple(row) for row in workbook[name].iter_rows(min_row=1, values_only=True)]
        names = ", ".join(f"«{name}»" for name in workbook.sheetnames)
        raise OntologyImportError(
            [f"The sheet «{SHEET_NAME}» is missing; the workbook has {names}."]
        )
    finally:
        workbook.close()


def _find_header(rows: Sequence[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    """Return the 0-based index of the header row and the column index of each field."""
    id_header = normalize_arabic(HEADERS["id"])
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = [normalize_arabic(cell) if isinstance(cell, str) else "" for cell in row]
        if id_header not in names:
            continue
        columns: dict[str, int] = {}
        problems = []
        for field_name, header in HEADERS.items():
            wanted = normalize_arabic(header)
            if names.count(wanted) == 1:
                columns[field_name] = names.index(wanted)
            else:
                problem = "is missing" if wanted not in names else "appears more than once"
                problems.append(f"Header row {index + 1}: the column «{header}» {problem}.")
        if problems:
            raise OntologyImportError(problems)
        return index, columns
    message = (
        f"No header row found in the first {HEADER_SEARCH_ROWS} rows: "
        f"a row holding the column «{HEADERS['id']}» is expected."
    )
    raise OntologyImportError([message])


# ─── One row ───


def _is_blank(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _valid_id(entity_id: str) -> bool:
    digits = entity_id[1:]
    return len(entity_id) >= 4 and entity_id[0] == "E" and digits.isascii() and digits.isdigit()


class _RowReader:
    """Reads the cells of one row and records what is wrong with them."""

    def __init__(self, number: int, row: tuple[Any, ...], columns: dict[str, int]) -> None:
        self.number = number
        self.row = row
        self.columns = columns
        self.problems: list[str] = []

    def raw_value(self, field_name: str) -> Any:
        index = self.columns[field_name]
        return self.row[index] if index < len(self.row) else None

    def report(self, field_name: str, message: str) -> None:
        index = self.columns[field_name]
        column = f"«{HEADERS[field_name]}» ({chr(ord('A') + index)})"
        self.problems.append(f"Row {self.number}, column {column}: {message}.")

    def text(self, field_name: str) -> str | None:
        """Return the cell as text; None when it is blank, or when it is not text (reported)."""
        value = self.raw_value(field_name)
        if _is_blank(value):
            return None
        if not isinstance(value, str):
            self.report(field_name, f"holds a value of type {type(value).__name__}, expected text")
            return None
        return value

    def entity_id(self) -> str:
        value = (self.text("id") or "").strip()
        if value and not _valid_id(value):
            self.report("id", f"«{value}» is not an id of the form E001")
        return value

    def items(self, field_name: str) -> tuple[str, ...]:
        cell = self.text(field_name)
        items = split_list(cell or "")
        if cell is not None and not items:
            self.report(field_name, "holds no item")
        return tuple(items)

    def label(self) -> str | None:
        value = self.text("label_ar")
        if value is not None and not normalize_arabic(value):
            self.report("label_ar", "has no letter")
        return value

    def constraint(self) -> tuple[str | None, ConstraintKind | None]:
        value = self.text("special_constraint")
        kind = classify_constraint(value)
        if kind is ConstraintKind.UNKNOWN:
            self.report(
                "special_constraint",
                f"«{value}» is a constraint the code cannot enforce; add it to "
                "ontology_constraints.KNOWN_CONSTRAINTS, with its behaviour, before importing",
            )
        return value, kind

    def empty_required_cells(self) -> None:
        for name in REQUIRED_FIELDS:
            if _is_blank(self.raw_value(name)):
                self.report(name, "is empty")


def _read_row(
    number: int, row: tuple[Any, ...], columns: dict[str, int]
) -> tuple[OntologyRow | None, list[str]]:
    """Return the entity of a row, or the problems that stop it from being one."""
    reader = _RowReader(number, row, columns)
    reader.empty_required_cells()
    entity_id = reader.entity_id()
    label = reader.label()
    lists = {name: reader.items(name) for name in LIST_FIELDS}
    constraint, kind = reader.constraint()
    domain = reader.text("domain")
    if reader.problems or label is None or domain is None:
        return None, reader.problems
    entity = OntologyRow(
        id=entity_id,
        label_ar=label,
        related_objects=lists["related_objects"],
        actions_and_uses=lists["actions_and_uses"],
        contextual_concepts=lists["contextual_concepts"],
        special_constraint=constraint,
        domain=domain,
        is_catch_all=kind is ConstraintKind.SPECIFY_BEFORE_SEARCH,
        raw={"row": number, **{name: reader.raw_value(name) for name in HEADERS}},
    )
    return entity, []


# ─── The whole sheet ───


def _id_number(entity_id: str) -> int:
    return int(entity_id[1:])


def _entity_id(number: int) -> str:
    return f"E{number:03d}"


def _id_list(ids: Sequence[str]) -> str:
    shown = ", ".join(ids[:MAX_IDS_SHOWN])
    return shown if len(ids) <= MAX_IDS_SHOWN else f"{shown}, ... ({len(ids)} in all)"


def _check_ids(
    rows: Sequence[OntologyRow], rows_of: dict[str, list[int]], expected: int
) -> list[str]:
    problems = [
        f"The id {entity_id} is used by {len(numbers)} rows: {', '.join(map(str, numbers))}."
        for entity_id, numbers in rows_of.items()
        if len(numbers) > 1
    ]
    if len(rows) != expected:
        problems.append(f"The sheet has {len(rows)} entities, {expected} were expected.")
    first, last = _entity_id(1), _entity_id(expected)
    wanted = {_entity_id(number) for number in range(1, expected + 1)}
    found = {row.id for row in rows}
    if missing := sorted(wanted - found, key=_id_number):
        problems.append(f"Ids missing between {first} and {last}: {_id_list(missing)}.")
    if unexpected := sorted(found - wanted, key=_id_number):
        problems.append(f"Ids outside {first} to {last}: {_id_list(unexpected)}.")
    return problems


def _duplicate_label_warnings(rows: Sequence[OntologyRow]) -> list[str]:
    seen: dict[str, str] = {}
    warnings = []
    for row in rows:
        key = normalize_arabic(row.label_ar)
        if key in seen:
            warnings.append(f"{row.id} has the same label as {seen[key]}: «{row.label_ar}».")
        else:
            seen[key] = row.id
    return warnings


def parse_workbook(
    data: bytes, *, source_name: str, expected_count: int = EXPECTED_COUNT
) -> ParsedOntology:
    """
    Validate the workbook bytes and return its entities.

    Checks the file, the sheet, the header and its seven columns, each cell
    (filled, text, an id of the form E001), the ids (unique, exactly E001 to
    E<expected_count>), the row count and the constraint texts. Two entities with
    the same label are not an error, but are reported in `warnings`.
    """
    rows = _read_sheet(data)
    header_index, columns = _find_header(rows)

    parsed: list[OntologyRow] = []
    problems: list[str] = []
    rows_of: dict[str, list[int]] = {}
    for number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if all(_is_blank(cell) for cell in row):
            continue
        entity, row_problems = _read_row(number, row, columns)
        problems.extend(row_problems)
        if entity is not None:
            parsed.append(entity)
            rows_of.setdefault(entity.id, []).append(number)
    problems.extend(_check_ids(parsed, rows_of, expected_count))
    if problems:
        raise OntologyImportError(problems)

    return ParsedOntology(
        source_name=source_name,
        source_sha256=hashlib.sha256(data).hexdigest(),
        sheet=SHEET_NAME,
        rows=tuple(sorted(parsed, key=lambda row: _id_number(row.id))),
        warnings=tuple(_duplicate_label_warnings(parsed)),
    )


def read_workbook(path: Path, *, expected_count: int = EXPECTED_COUNT) -> ParsedOntology:
    """Read the file at `path` and validate it; a file that cannot be read is an import error too."""
    try:
        data = path.read_bytes()
    except OSError as error:
        reason = error.strerror or type(error).__name__
        raise OntologyImportError([f"Cannot read {path}: {reason}."]) from None
    return parse_workbook(data, source_name=path.name, expected_count=expected_count)


# ─── The generated JSON ───


def _entity_values(row: OntologyRow) -> dict[str, Any]:
    return {
        "id": row.id,
        "label_ar": row.label_ar,
        "related_objects": list(row.related_objects),
        "actions_and_uses": list(row.actions_and_uses),
        "contextual_concepts": list(row.contextual_concepts),
        "special_constraint": row.special_constraint,
        "domain": row.domain,
        "is_catch_all": row.is_catch_all,
        "raw": row.raw,
    }


def export_json(parsed: ParsedOntology) -> str:
    """
    Return the generated JSON: the source (file name, hash, count) and the entities.

    The text has no timestamp, so the same workbook always gives the same bytes and
    a regenerated file shows in `git diff` only when the workbook changed. Entities
    are one per line.
    """
    source = {
        "file": parsed.source_name,
        "sha256": parsed.source_sha256,
        "sheet": parsed.sheet,
        "entity_count": len(parsed.rows),
        "domains": parsed.domains,
        "constraint_counts": parsed.constraint_counts,
        "generated_by": "apps/api: python -m src.cli.import_ontology",
    }
    source_text = json.dumps(source, ensure_ascii=False, indent=2).replace("\n", "\n  ")
    entities = ",\n".join(
        "    " + json.dumps(_entity_values(row), ensure_ascii=False) for row in parsed.rows
    )
    return f'{{\n  "source": {source_text},\n  "entities": [\n{entities}\n  ]\n}}\n'


# ─── The database ───


def search_forms(row: OntologyRow) -> tuple[str, list[str], str]:
    """Return the label, the related objects and the whole text in search form."""
    label = normalize_arabic(row.label_ar)
    related = [term for item in row.related_objects if (term := normalize_arabic(item))]
    related = list(dict.fromkeys(related))
    return label, related, " ".join(dict.fromkeys([label, *related]))


def _row_values(row: OntologyRow, source_sha256: str) -> dict[str, Any]:
    label_norm, related_norm, search_text = search_forms(row)
    return {
        **_entity_values(row),
        "label_norm": label_norm,
        "related_norm": related_norm,
        "search_text": search_text,
        "source_sha256": source_sha256,
    }


def _batches(values: list[dict[str, Any]]) -> Iterator[list[dict[str, Any]]]:
    for batch in batched(values, ROWS_PER_INSERT):
        yield list(batch)


async def load_ontology(session: AsyncSession, parsed: ParsedOntology) -> LoadResult:
    """
    Make `app.ontology_entities` hold exactly the entities of `parsed`.

    Entities are inserted or updated by id, and the ones the workbook no longer
    has are removed (a candidate that was folded into one keeps its row). The
    caller owns the transaction: nothing is committed here.
    """
    values = [_row_values(row, parsed.source_sha256) for row in parsed.rows]
    inserted = updated = 0
    for batch in _batches(values):
        insertion = insert(OntologyEntity).values(batch)
        replaced = {name: insertion.excluded[name] for name in batch[0] if name != "id"}
        upsert = insertion.on_conflict_do_update(
            index_elements=[OntologyEntity.id], set_={**replaced, "imported_at": func.now()}
        ).returning(literal_column("xmax = 0", Boolean))
        was_inserted = (await session.scalars(upsert)).all()
        inserted += sum(was_inserted)
        updated += len(was_inserted) - sum(was_inserted)

    removed = await session.scalars(
        delete(OntologyEntity)
        .where(OntologyEntity.id.not_in([row.id for row in parsed.rows]))
        .returning(OntologyEntity.id)
    )
    return LoadResult(inserted=inserted, updated=updated, removed=len(removed.all()))
